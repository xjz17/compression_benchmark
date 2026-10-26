#!/usr/bin/env python3
"""Batch-enrich the reading-list spreadsheet with Semantic Scholar + OpenAlex metadata.

Reads `benchmark/reference/reading list - compress.xlsx` (sheet ``paper``) and writes a
review CSV with one row per paper: S2 hit quality, identifiers, and any ``github.com`` URLs
scraped from the Semantic Scholar *abstract* (when authors paste links there).

Papers With Code REST (official code links) is optionally called when ``--try-pwc`` is set;
it may fail on some Python/LibreSSL installs (TLS handshake). Prefer a normal OpenSSL/Python
environment or skip PWC and rely on S2 abstract + manual review.

Environment:
  S2_API_KEY   Semantic Scholar API key (strongly recommended for 300+ rows; higher limits).

Usage (from repo root)::

  python3 scripts/fetch_reading_list_code_hints.py \\
      --output benchmark/reference/reading_list_code_audit.csv

  # Resume after interrupt: merges with existing output by Title+File key.
  python3 scripts/fetch_reading_list_code_hints.py --resume

Deps: ``openpyxl`` (``pip install openpyxl`` in a venv).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import random
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "benchmark" / "reference" / "reading list - compress.xlsx"
DEFAULT_OUT = ROOT / "benchmark" / "reference" / "reading_list_code_audit.csv"
DEFAULT_CACHE = ROOT / "benchmark" / "reference" / ".reading_list_api_cache.json"

GH_IN_TEXT = re.compile(
    r"https://github\.com/([A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+)(?:/[^\s\])\"'<>]*)?",
    re.IGNORECASE,
)


def _norm_repo(url: str) -> str:
    m = GH_IN_TEXT.search(url)
    if not m:
        return ""
    path = m.group(1)
    if "/" not in path:
        return ""
    owner, _, repo = path.partition("/")
    # Abstracts often end a URL with a sentence period: github.com/org/Tool.
    owner, repo = owner.rstrip("."), repo.rstrip(".")
    if not owner or not repo:
        return ""
    return f"https://github.com/{owner}/{repo}"


def _short_authors(authors: list[dict[str, Any]] | None, limit: int = 3) -> str:
    if not authors:
        return ""
    names = [a.get("name") or "" for a in authors[:limit]]
    s = ", ".join(n for n in names if n)
    if len(authors) > limit:
        s += ", …"
    return s


def _http_json(
    url: str,
    *,
    headers: dict[str, str],
    timeout: float = 60.0,
) -> tuple[int, Any]:
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            body = resp.read().decode("utf-8", errors="replace")
            return resp.status, json.loads(body)
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8", errors="replace")
        try:
            return e.code, json.loads(raw)
        except json.JSONDecodeError:
            return e.code, {"_raw": raw}


def ss_search_first(
    title: str,
    *,
    api_key: str | None,
    fields: str,
    sleep_s: float,
    max_429_retries: int = 4,
) -> tuple[dict[str, Any] | None, str | None]:
    if not title.strip():
        return None, "empty_title"
    q = urllib.parse.quote(title.strip())
    base = f"https://api.semanticscholar.org/graph/v1/paper/search?query={q}&limit=3&fields={fields}"
    headers = {"User-Agent": "web_compression-reading-list-script/1.0 (research audit)"}
    if api_key:
        headers["x-api-key"] = api_key
    backoff = 12.0
    for attempt in range(max_429_retries + 1):
        time.sleep(sleep_s + random.uniform(0, 0.35))
        status, data = _http_json(base, headers=headers)
        if status == 429 and attempt < max_429_retries:
            # Public tier is very small; exponential backoff before giving up.
            time.sleep(backoff + random.uniform(0, 2.0))
            backoff *= 1.8
            continue
        if status == 429:
            return None, "s2_429_rate_limited"
        if status != 200:
            return None, f"s2_http_{status}:{data!s}"[:500]
        hits = data.get("data") or []
        if not hits:
            return None, "s2_no_results"
        return hits[0], None
    return None, "s2_429_rate_limited"


def openalex_first_title(title: str, *, mailto: str, sleep_s: float) -> tuple[dict[str, Any] | None, str | None]:
    if not title.strip():
        return None, "empty_title"
    params = urllib.parse.urlencode({"search": title.strip(), "per_page": 1})
    url = f"https://api.openalex.org/works?{params}"
    headers = {"User-Agent": f"web_compression-reading-list-script/1.0 ({mailto})"}
    time.sleep(sleep_s + random.uniform(0, 0.25))
    status, data = _http_json(url, headers=headers)
    if status != 200:
        return None, f"openalex_http_{status}"
    results = data.get("results") or []
    if not results:
        return None, "openalex_no_results"
    return results[0], None


def pwc_by_arxiv(arxiv_id: str, *, sleep_s: float) -> tuple[list[str], str | None]:
    """Return list of GitHub repo HTTPS URLs from Papers with Code, if API works."""
    aid = arxiv_id.strip().replace("arXiv:", "").replace("arxiv:", "")
    if not aid:
        return [], None
    qs = urllib.parse.urlencode({"arxiv_id": aid, "items_per_page": 5})
    url = f"https://api.paperswithcode.com/v1/papers/?{qs}"
    headers = {"User-Agent": "web_compression-reading-list-script/1.0 (optional PWC)"}
    time.sleep(sleep_s)
    try:
        status, data = _http_json(url, headers=headers)
    except urllib.error.URLError as e:
        return [], f"pwc_network:{e!s}"[:300]
    if status != 200:
        return [], f"pwc_http_{status}"
    repos: list[str] = []

    def collect(u: Any) -> None:
        if isinstance(u, str) and "github.com" in u.lower():
            nr = _norm_repo(u)
            if nr:
                repos.append(nr)

    for paper in data.get("results") or []:
        for r in paper.get("repository_url_list") or []:
            collect(r)
        for repo in paper.get("repositories") or []:
            if isinstance(repo, dict):
                collect(repo.get("url"))
        repo_one = paper.get("repository")
        if isinstance(repo_one, dict):
            collect(repo_one.get("url"))
    return sorted(set(repos)), None


def load_cache(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def save_cache(path: Path, cache: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def cache_key(prefix: str, title: str, file_key: str) -> str:
    h = hashlib.sha256(f"{prefix}|{title}|{file_key}".encode("utf-8")).hexdigest()[:24]
    return f"{prefix}:{h}"


def read_excel_rows(xlsx: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as e:
        raise SystemExit(
            "openpyxl is required: python3 -m venv .venv && . .venv/bin/activate && pip install openpyxl"
        ) from e
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "paper" not in wb.sheetnames:
        raise SystemExit(f"No sheet named 'paper' in {xlsx}; found {wb.sheetnames}")
    ws = wb["paper"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: i for i, name in enumerate(header) if name}
    rows: list[dict[str, Any]] = []
    for r_i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def get(name: str) -> str:
            i = col.get(name)
            if i is None or i >= len(row):
                return ""
            v = row[i]
            if v is None:
                return ""
            return str(v).strip()

        title = get("Title")
        if not title:
            continue
        rows.append(
            {
                "excel_row": r_i,
                "Name": get("Name"),
                "File": get("File"),
                "Title": title,
                "Author": get("Author"),
                "Institute": get("Institute"),
                "excel_Code": get("Code"),
            }
        )
    wb.close()
    return rows


def write_audit_csv(path: Path, rows: list[dict[str, Any]], field_order: list[str] | None) -> list[str]:
    """Rewrite the full audit CSV (crash-safe incremental runs). Returns updated column order."""
    if not rows:
        return field_order or []
    keys = list(field_order) if field_order else []
    for r in rows:
        for k in r:
            if k not in keys:
                keys.append(k)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in keys})
    return keys


def load_existing_output(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    if not path.is_file():
        return {}
    out: dict[tuple[str, str], dict[str, str]] = {}
    with path.open(newline="", encoding="utf-8") as f:
        r = csv.DictReader(f)
        for row in r:
            t = (row.get("Title") or "", row.get("File") or "")
            if t[0]:
                out[t] = {k: row.get(k, "") for k in r.fieldnames or []}
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="Enrich reading list with S2/OpenAlex + GitHub hints.")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--output", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--cache", type=Path, default=DEFAULT_CACHE)
    ap.add_argument("--mailto", default="https://github.com/xiaojinzhao/web_compression")
    ap.add_argument("--sleep-ss", type=float, default=3.0, help="Base delay between S2 calls (no key).")
    ap.add_argument(
        "--sleep-ss-key",
        type=float,
        default=1.0,
        help="Base delay when S2_API_KEY is set (tune to stay under your quota).",
    )
    ap.add_argument("--sleep-openalex", type=float, default=0.35)
    ap.add_argument("--sleep-pwc", type=float, default=0.25)
    ap.add_argument("--try-pwc", action="store_true", help="Query Papers with Code when ArXiv id exists.")
    ap.add_argument("--resume", action="store_true", help="Skip Title+File rows already present in --output.")
    ap.add_argument(
        "--max-rows",
        type=int,
        default=0,
        help="Only read the first N spreadsheet rows with titles (0 = all; for dry runs).",
    )
    args = ap.parse_args()

    api_key = os.environ.get("S2_API_KEY") or os.environ.get("SEMANTIC_SCHOLAR_API_KEY")
    ss_sleep = args.sleep_ss_key if api_key else args.sleep_ss
    fields = ",".join(
        [
            "paperId",
            "title",
            "year",
            "authors",
            "externalIds",
            "url",
            "openAccessPdf",
            "abstract",
        ]
    )

    rows = read_excel_rows(args.xlsx)
    if args.max_rows:
        rows = rows[: args.max_rows]
    existing = load_existing_output(args.output) if args.resume else {}
    cache = load_cache(args.cache)

    out_rows: list[dict[str, Any]] = []
    field_order: list[str] | None = None
    for rec in rows:
        title = rec["Title"]
        file_key = rec["File"]
        k = (title, file_key)
        if args.resume and k in existing and existing[k].get("s2_paper_id"):
            out_rows.append(existing[k])
            field_order = write_audit_csv(args.output, out_rows, field_order)
            print(f"[{len(out_rows)}/{len(rows)}] resume  {title[:72]}…", flush=True)
            continue

        ck_ss = cache_key("ss", title, file_key)
        cached_ss = cache.get(ck_ss) if ck_ss in cache else None
        if cached_ss:
            hit, err = cached_ss.get("hit"), cached_ss.get("error")
        else:
            hit, err = ss_search_first(title, api_key=api_key, fields=fields, sleep_s=ss_sleep)
            cache[ck_ss] = {"hit": hit, "error": err, "ts": time.time()}
            save_cache(args.cache, cache)

        abstract_gh: list[str] = []
        if hit and hit.get("abstract"):
            abstract_gh = sorted({_norm_repo(m.group(0)) for m in GH_IN_TEXT.finditer(hit["abstract"])})
        ext = (hit or {}).get("externalIds") or {}
        arxiv = ext.get("ArXiv") or ""

        pwc_gh: list[str] = []
        pwc_err = ""
        if args.try_pwc and arxiv:
            pwc_gh, pwc_err = pwc_by_arxiv(arxiv, sleep_s=args.sleep_pwc)

        ck_oa = cache_key("oa", title, file_key)
        if ck_oa in cache:
            oa_hit, oa_err = cache[ck_oa].get("hit"), cache[ck_oa].get("error")
        else:
            oa_hit, oa_err = openalex_first_title(title, mailto=args.mailto, sleep_s=args.sleep_openalex)
            cache[ck_oa] = {"hit": oa_hit, "error": oa_err, "ts": time.time()}
            save_cache(args.cache, cache)

        oa_id = (oa_hit or {}).get("id") or ""
        oa_doi = (oa_hit or {}).get("doi") or ""

        merged_gh = sorted(set(abstract_gh) | set(pwc_gh))

        out = {
            **rec,
            "s2_paper_id": (hit or {}).get("paperId") or "",
            "s2_title": (hit or {}).get("title") or "",
            "s2_year": (hit or {}).get("year") or "",
            "s2_authors_short": _short_authors((hit or {}).get("authors")),
            "s2_url": (hit or {}).get("url") or "",
            "s2_doi": ext.get("DOI") or "",
            "s2_arxiv": arxiv,
            "s2_open_pdf": ((hit or {}).get("openAccessPdf") or {}).get("url") or "",
            "s2_error": err or "",
            "github_candidates_ss_abstract": "; ".join(abstract_gh),
            "github_candidates_pwc": "; ".join(pwc_gh),
            "pwc_error": pwc_err,
            "github_candidates_merged": "; ".join(merged_gh),
            "openalex_id": oa_id,
            "openalex_doi": oa_doi,
            "openalex_url": oa_id,
            "openalex_error": oa_err or "",
            "manual_verified_repo": "",
            "manual_notes": "",
        }
        out_rows.append(out)
        field_order = write_audit_csv(args.output, out_rows, field_order)
        hint = "OK" if (hit and not err) else (err or "no_hit")
        print(f"[{len(out_rows)}/{len(rows)}] {hint}  {title[:72]}…", flush=True)

    if not out_rows and not rows:
        print("No rows with Title in spreadsheet.", file=sys.stderr)
        return

    field_order = write_audit_csv(args.output, out_rows, field_order)

    missing_key = "" if api_key else " (set S2_API_KEY for faster/higher quota)"
    print(f"Wrote {len(out_rows)} rows to {args.output}{missing_key}", flush=True)


if __name__ == "__main__":
    main()
